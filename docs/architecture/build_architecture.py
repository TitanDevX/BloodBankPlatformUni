from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
OUT = ROOT.parent / 'Software_Architecture.docx'
FONT = '/System/Library/Fonts/HelveticaNeue.ttc'

sources = {
'context.mmd': '''flowchart LR
  A[Administrator] --> P[Blood Bank Platform]
  B[Blood bank staff] --> P
  L[Laboratory staff] --> P
  H[Hospital staff] --> P
  P --> E[Email provider]
''',
'application.mmd': '''flowchart LR
  U[Staff browser] --> R[React frontend]
  R -->|HTTPS JSON| L[Laravel modular monolith]
  L -->|Queries and transactions| D[(Relational database)]
  L -->|HTTP JSON prepared demand history| P[Python forecasting service]
  P --> M[Saved model]
  P -->|Predictions| L
''',
'modules.mmd': '''flowchart TB
  subgraph Laravel modules
    I[Identity]
    O[Organizations]
    D[Donors]
    N[Donations]
    S[Screening]
    V[Inventory]
    H[HospitalRequests]
    R[Recommendations]
    F[Forecasting]
    T[Traceability]
    P[Reports]
    C[Core]
  end
''',
'deployment.mmd': '''flowchart LR
  B[Staff browser React] -->|HTTPS| W[Demonstration web host Laravel API]
  W -->|Private database connection| D[(Relational database)]
  W -->|HTTP JSON| P[Python forecasting runtime]
  P --> M[Saved model file]
'''
}
for name, value in sources.items():
    (ROOT / name).write_text(value)

def font(size, bold=False):
    return ImageFont.truetype(FONT, size, index=1 if bold else 0)

def box(d, xy, lines, fill='#F3F5F7', outline='#A5B2BA', size=28):
    d.rounded_rectangle(xy, radius=14, fill=fill, outline=outline, width=2)
    f=font(size, True)
    text='\n'.join(lines)
    bb=d.multiline_textbbox((0,0),text,font=f,spacing=7,align='center')
    x=(xy[0]+xy[2]-(bb[2]-bb[0]))/2
    y=(xy[1]+xy[3]-(bb[3]-bb[1]))/2
    d.multiline_text((x,y),text,font=f,fill='#1D2B34',spacing=7,align='center')

def arrow(d, start, end, label=None, label_xy=None):
    import math
    d.line((start,end),fill='#596D78',width=4)
    angle=math.atan2(end[1]-start[1],end[0]-start[0]); a=17
    p1=(end[0]-a*math.cos(angle-.5),end[1]-a*math.sin(angle-.5))
    p2=(end[0]-a*math.cos(angle+.5),end[1]-a*math.sin(angle+.5))
    d.polygon([end,p1,p2],fill='#596D78')
    if label and label_xy:
        d.text(label_xy,label,font=font(18),fill='#354650')

def canvas(name, w, h):
    im=Image.new('RGB',(w,h),'white'); return im,ImageDraw.Draw(im)

im,d=canvas('context',1400,640)
for y,label in [(60,'Administrator'),(190,'Blood bank staff'),(320,'Laboratory staff'),(450,'Hospital staff')]:
    box(d,(40,y,385,y+92),[label],size=25)
    arrow(d,(385,y+46),(560,305))
box(d,(560,225,960,385),['Blood Bank Platform'],fill='#E8F0F3',size=32)
box(d,(1100,255,1360,355),['Email provider'],size=25)
arrow(d,(960,305),(1100,305),'follow up email',(965,268))
im.save(ROOT/'context.png')

im,d=canvas('application',1600,560)
items=[(25,175,265,355,['Staff browser']),(345,175,600,355,['React frontend']),(690,175,1030,355,['Laravel API','modules']),(1140,35,1550,195,['Relational database']),(1140,345,1550,515,['Python forecast','service and model'])]
for x1,y1,x2,y2,l in items: box(d,(x1,y1,x2,y2),l,fill='#E8F0F3' if l[0].startswith('Laravel') else '#F3F5F7')
arrow(d,(265,265),(345,265)); arrow(d,(600,265),(690,265),'HTTPS / JSON',(595,220))
arrow(d,(1030,225),(1140,115),'queries',(1030,145)); arrow(d,(1030,315),(1140,430),'HTTP / JSON',(1035,380))
im.save(ROOT/'application.png')

im,d=canvas('modules',1500,640)
modules=[('Identity',40,65),('Organizations',350,65),('Donors',660,65),('Donations',970,65),('Screening',1280,65),('Inventory',40,250),('HospitalRequests',350,250),('Recommendations',660,250),('Forecasting',970,250),('Traceability',1280,250),('Reports',505,435),('Core',850,435)]
for label,x,y in modules:
    long=label in ('HospitalRequests','Recommendations')
    box(d,(x,y,x+(280 if long else 190),y+115),[label],fill='#E8F0F3' if label in ('Inventory','HospitalRequests') else '#F3F5F7',size=19 if long else 22)
d.text((225,585),'Modules own business changes; contracts and services coordinate cross module work.',font=font(24),fill='#354650')
im.save(ROOT/'modules.png')

im,d=canvas('deployment',1500,560)
box(d,(40,170,320,345),['Staff browser','React client'],size=25)
box(d,(430,170,850,345),['Demonstration web host','Laravel API'],fill='#E8F0F3',size=28)
box(d,(1000,45,1460,200),['Relational database'],size=26)
box(d,(1000,330,1460,500),['Python runtime','saved model'],size=26)
arrow(d,(320,255),(430,255),'HTTPS',(345,210))
arrow(d,(850,220),(1000,125),'private connection',(853,145))
arrow(d,(850,295),(1000,410),'HTTP / JSON',(860,365))
im.save(ROOT/'deployment.png')

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.68)
sec.left_margin=Inches(.82); sec.right_margin=Inches(.82)

styles=doc.styles
normal=styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor(25,34,40)
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.12
for s,size,before,after in [('Title',22,0,12),('Heading 1',14,15,7),('Heading 2',11.5,10,5)]:
    st=styles[s]; st.font.name='Aptos Display' if s=='Title' else 'Aptos'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.keep_with_next=True
    if s=='Title':
        ppr=st.element.get_or_add_pPr()
        for b in ppr.findall(qn('w:pBdr')): ppr.remove(b)
caption=styles.add_style('Figure Caption',WD_STYLE_TYPE.PARAGRAPH)
caption.font.name='Aptos'; caption.font.size=Pt(8.5); caption.font.italic=True; caption.font.color.rgb=RGBColor(75,85,92)
caption.paragraph_format.space_after=Pt(10); caption.paragraph_format.keep_with_next=False

def para(t,style=None): return doc.add_paragraph(t,style)
def heading(t,l=1): doc.add_heading(t,l)
def figure(filename, caption_text, width=6.75):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(ROOT/filename),width=Inches(width))
    doc.add_paragraph(caption_text,'Figure Caption')
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers): t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=v
    for ri,row in enumerate(t.rows):
        for i,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths: c.width=Inches(widths[i])
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
                for r in p.runs: r.font.size=Pt(8.5)
            if ri==0:
                tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'DDE7EA'); tcPr.append(sh)
                for p in c.paragraphs:
                    for r in p.runs: r.bold=True
    trPr=t.rows[0]._tr.get_or_add_trPr()
    repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); trPr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

doc.add_paragraph('Software Architecture','Title')
para('Blood Bank Platform  |  Graduation 1 project report')
para('This chapter describes the intended structure of the Blood Bank Platform and the boundaries that keep its clinical records, stock allocation and forecasting work consistent. It follows the approved Graduation 1 requirements. The repository currently contains a Laravel scaffold and design artefacts; the application modules, React client and Python service shown here are planned components.')

heading('Architecture overview')
para('The platform is designed as a modular monolith. One Laravel application will expose the API and group backend code by business responsibility. A React client will call that API, while a separate Python service will calculate demand forecasts from prepared request history. Laravel modules will share one relational database. This arrangement lets the two application developers work against explicit module boundaries without operating several business services; Baraa can develop and evaluate the forecast service independently.')

heading('System context')
figure('context.png','Figure 1. Users and external email service at the platform boundary.')
para('Administrators maintain participating organizations and staff accounts. Blood bank staff register donors and donations, manage components, confirm reservations and record issue. Laboratory staff enter screening results and release decisions. Hospital staff create and follow requests. Donors have staff-managed records but no software login. The email provider supports screening follow-up; collection work and laboratory equipment remain outside the software boundary.')

heading('Application structure')
figure('application.png','Figure 2. Planned application components and data paths.')
para('The browser sends authenticated JSON requests to Laravel over HTTPS. Laravel validates inputs, enforces role and organization scope, and owns database changes. The Forecasting module sends prepared historical requested-unit demand to Python and receives predictions and model metadata. The Python service uses a saved model and does not require unrestricted database access. Laravel stores forecast runs and points for later viewing; a forecast cannot approve or allocate stock.')

heading('Backend modules and ownership')
figure('modules.png','Figure 3. Proposed Laravel module boundaries. Lines are omitted because modules communicate through named services and contracts rather than a fixed dependency chain.')
table(['Module','Main responsibility'],[
('Organizations','Banks, hospitals, active status and coordinates'),('Identity','Users, authentication, roles and organization scope'),('Donors','Staff-managed donor profiles and history'),('Donations','Eligibility, collection and processing records'),('Screening','Results, release decisions and donor follow-up'),('Inventory','Components, availability, reservations, expiry and disposal'),('HospitalRequests','Patient-linked requests, status and issuing workflow'),('Recommendations','Exact filtering and one nearest qualifying source'),('Forecasting','Python integration, runs and points'),('Traceability','Record-chain lookup and action history'),('Reports','Operational summaries and exports'),('Core','Small shared support utilities')], [1.55,5.2])
para('Each module owns changes to its records and exposes defined services for other modules. Controllers do not bypass those services to update another module’s tables. Permission and organization checks remain on the Laravel side. Confirmation, reservation and issue each use a database transaction to keep request and component state aligned. Core is limited to genuinely shared support code.')

heading('Inside a module')
para('Inventory is a representative module. Its routes lead to controllers; request classes validate input; policies check the acting user; services apply the stock rules; models persist changes; resources shape API responses. Contracts define the operations other modules may call. Module migrations and tests stay beside the code they support. The same pattern can be reused without forcing every module to have identical files.')
p=doc.add_paragraph()
p.paragraph_format.left_indent=Inches(.3)
p.paragraph_format.space_after=Pt(8)
run=p.add_run('Inventory/\n  app/Http/{Controllers, Requests, Resources}/\n  app/{Models, Services, Policies, Contracts, Providers}/\n  database/   routes/   tests/')
run.font.name='Menlo'; run.font.size=Pt(8)

heading('Runtime interactions and transaction boundaries')
table(['Interaction','Architectural rule'],[
('Login and access','Identity authenticates active staff; each protected operation checks role and organization scope.'),
('Source search','Recommendations reads eligible stock and returns only the nearest bank able to meet the whole exact-match request; search does not reserve stock.'),
('Confirm and reserve','Authorized source-bank staff confirm selected containers. Inventory locks and rechecks them, then writes the full reservation set and request status in one transaction or writes nothing.'),
('Issue and fulfill','Inventory and HospitalRequests recheck the complete reserved quantity and current eligibility. Issue records, component states and request fulfillment commit together.'),
('Generate forecast','Forecasting prepares historical requested-unit series, calls Python and saves the returned run and points. Failure or insufficient history is shown as unavailable.')],[1.65,5.1])
para('The existing login sequence diagram in docs/uml covers authentication detail. The remaining allocation, issue and forecast sequences should be drawn with their use-case documentation as those workflows are implemented.')

heading('Data architecture')
para('The planned relational database is shared by the Laravel modules. Bank and hospital identifiers establish record ownership. The main trace runs from donor to donation to component, then through reservation to hospital request and issue. Donation attempts remain recorded even when collection fails. Inactive organizations and accounts retain their historical links; significant actions retain actor and time information. ForecastRun and ForecastPoint store target dates, requested-unit predictions, model version and data provenance.')
para('A component represents one physical container and one allocatable unit. One bank must cover a request in full. To prevent double allocation, reservation confirmation should lock candidate component rows, recheck approval, exact specifications and expiry, and enforce at most one active reservation per component. Cancellation and invalidation release a reservation set atomically; only still eligible stock becomes available again. The master ERD is in docs/erd/main_erd.png, with editable source in docs/erd/mermaid_only/main_erd.mmd.')

heading('Deployment view')
figure('deployment.png','Figure 4. Planned demonstration deployment; host and database products remain to be selected.')
para('Local development currently consists of the Laravel project and a local SQLite file, with frontend and AI directories reserved for future work. The demonstration target is a browser-hosted React client, a Laravel API host, a relational database and a reachable Python forecasting runtime. HTTPS protects browser-to-API traffic. The API-to-Python connection carries prepared demand history and predictions. Final hostnames, database product, process layout and service security settings must be recorded after the demonstration environment is chosen.')

heading('Decisions and limits')
table(['Decision','Reason','Trade-off'],[
('Modular Laravel monolith','Clear ownership with one main backend','Boundaries need code review and discipline'),
('Separate Python service','Independent model development and evaluation','Service outages need an unavailable response'),
('Shared relational database','Atomic stock transactions and simple integration','Modules share database infrastructure'),
('React client','Reusable interactive screens','API integration adds testing work')],[1.45,2.7,2.6])
para('The prototype represents a simulated network. Source selection uses exact component, ABO/Rh and mandatory attribute matching, then straight-line distance with a stable bank-ID tie break. It offers no split-bank fulfillment, clinical substitution or automatic clinical decision. Forecast quality depends on documented historical or simulated request data. “Fulfilled” ends at source-bank issue; delivery and transfusion are outside Graduation 1. The fixed component shelf-life configuration is a prototype policy that must be agreed before demonstration.')

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Blood Bank Platform  |  Software Architecture').font.size=Pt(8)
doc.save(OUT)
print(OUT)
